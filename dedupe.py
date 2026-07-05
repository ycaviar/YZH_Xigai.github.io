# -*- coding: utf-8 -*-
"""
题库去重脚本：
1. 读取 questions.json，按标准化题干分组去重
2. 构建全局题池 pool（含 appearedIn 元数据）
3. 答案冲突组用多数投票定推荐答案，标记 conflict 并保留 answerVariants
4. 套题改为 questionUids 引用
5. 输出 questions_dedup.json
6. 导出 answer_conflicts.docx 供人工核查
"""
import json
import re
import os
from collections import defaultdict, Counter

SRC = r'd:\桌面\习概题库\questions.json'
OUT_JSON = r'd:\桌面\习概题库\questions_dedup.json'
OUT_DOCX = r'd:\桌面\习概题库\answer_conflicts.docx'


def normalize(s):
    if not s:
        return ''
    s = re.sub(r'\s+', '', s)
    s = re.sub(r'[，。、,．.\?？!！:：;；""''\'\"\(\)（）【】\[\]《》<>\-—_·]', '', s)
    return s.lower()


def type_label(t):
    return {'judge': '判断题', 'single': '单选题', 'multi': '多选题'}.get(t, t)


def majority_vote(answers):
    """返回出现次数最多的答案；平票时取列表中第一个。"""
    cnt = Counter(answers)
    max_count = max(cnt.values())
    for a in answers:  # 保持首次出现顺序
        if cnt[a] == max_count:
            return a
    return answers[0]


with open(SRC, encoding='utf-8') as f:
    data = json.load(f)

sets = data['sets']

# 收集所有题目（按套号、题号顺序）
all_items = []  # [{setId, qId, type, stem, stemNorm, options, answer}]
for s in sets:
    for q in s['questions']:
        all_items.append({
            'setId': s['id'],
            'qId': q['id'],
            'type': q['type'],
            'stem': q['stem'],
            'stemNorm': normalize(q['stem']),
            'options': q.get('options', []),
            'answer': q['answer']
        })

# 按标准化题干分组，保留首次出现顺序
groups = defaultdict(list)
first_seen_order = []
for item in all_items:
    key = item['stemNorm']
    if key not in groups:
        first_seen_order.append(key)
    groups[key].append(item)

# 构建全局题池
pool = []
stemnorm_to_uid = {}
for key in first_seen_order:
    items = groups[key]
    first = items[0]
    answers = [it['answer'] for it in items]
    distinct_answers = set(answers)
    is_conflict = len(distinct_answers) > 1
    recommended = majority_vote(answers)

    uid = 'q{:04d}'.format(len(pool) + 1)
    stemnorm_to_uid[key] = uid

    pool.append({
        'uid': uid,
        'type': first['type'],
        'stem': first['stem'],
        'options': first['options'],
        'answer': recommended,
        'conflict': is_conflict,
        'answerVariants': dict(Counter(answers)) if is_conflict else None,
        'appearedIn': [
            {'setId': it['setId'], 'qId': it['qId'], 'originalAnswer': it['answer']}
            for it in sorted(items, key=lambda x: (x['setId'], x['qId']))
        ]
    })

# 构建套题引用
new_sets = []
for s in sets:
    uids = []
    for q in s['questions']:
        uid = stemnorm_to_uid[normalize(q['stem'])]
        uids.append(uid)
    new_sets.append({
        'id': s['id'],
        'title': s['title'],
        'source': s['source'],
        'questionUids': uids
    })

# 统计
type_counts = Counter(p['type'] for p in pool)
conflict_count = sum(1 for p in pool if p['conflict'])
out = {
    'meta': {
        'total_unique': len(pool),
        'total_sets': len(new_sets),
        'total_references': len(all_items),
        'judge': type_counts.get('judge', 0),
        'single': type_counts.get('single', 0),
        'multi': type_counts.get('multi', 0),
        'conflicts': conflict_count
    },
    'pool': pool,
    'sets': new_sets
}

with open(OUT_JSON, 'w', encoding='utf-8') as f:
    json.dump(out, f, ensure_ascii=False, indent=2)

print(f"已写出: {OUT_JSON}")
print(f"唯一题数: {len(pool)} (原 {len(all_items)} 条, 去重率 {1 - len(pool)/len(all_items):.1%})")
print(f"套数: {len(new_sets)}")
print(f"答案冲突组: {conflict_count}")

# ============ 导出冲突到 docx ============
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = Document()

# 标题
h = doc.add_heading('题库答案冲突待核查清单', level=0)
h.alignment = WD_ALIGN_PARAGRAPH.CENTER

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run(f'共 {conflict_count} 组冲突 · 来源: questions.json (去重分析)')
run.font.size = Pt(11)
run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

doc.add_paragraph()  # 空行

conflict_pool = [p for p in pool if p['conflict']]

for idx, pq in enumerate(conflict_pool, 1):
    # 题号 + 题型
    h2 = doc.add_heading('', level=2)
    run = h2.add_run(f'冲突 #{idx}  [{type_label(pq["type"])}]  推荐答案: {pq["answer"]}')
    run.font.size = Pt(13)

    # 题干
    p_stem = doc.add_paragraph()
    r1 = p_stem.add_run('题干: ')
    r1.bold = True
    p_stem.add_run(pq['stem'])

    # 选项（非判断题）
    if pq['options']:
        p_opt = doc.add_paragraph()
        r2 = p_opt.add_run('选项:')
        r2.bold = True
        for opt in pq['options']:
            doc.add_paragraph(f'{opt["key"]}. {opt["text"]}', style='List Bullet')

    # 各套出现情况表格
    table = doc.add_table(rows=1, cols=4)
    table.style = 'Light Grid Accent 1'
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    hdr = table.rows[0].cells
    hdr[0].text = '套号'
    hdr[1].text = '原题号'
    hdr[2].text = '原答案'
    hdr[3].text = '与推荐一致?'
    for cell in hdr:
        for para in cell.paragraphs:
            for r in para.runs:
                r.bold = True

    for occ in pq['appearedIn']:
        row = table.add_row().cells
        row[0].text = f'第{occ["setId"]}套'
        row[1].text = str(occ['qId'])
        row[2].text = occ['originalAnswer']
        agree = '是' if occ['originalAnswer'] == pq['answer'] else '否 ⚠'
        row[3].text = agree
        if occ['originalAnswer'] != pq['answer']:
            for para in row[3].paragraphs:
                for r in para.runs:
                    r.font.color.rgb = RGBColor(0xcc, 0x00, 0x00)
                    r.bold = True

    # 答案分布
    p_var = doc.add_paragraph()
    r3 = p_var.add_run('答案分布: ')
    r3.bold = True
    variants_str = '  ·  '.join(f'"{k}" 出现 {v} 次' for k, v in pq['answerVariants'].items())
    p_var.add_run(variants_str)

    # 标记列（供人工填写）
    p_mark = doc.add_paragraph()
    r4 = p_mark.add_run('人工裁决: [   ]')
    r4.font.size = Pt(11)

    doc.add_paragraph()  # 组间空行

doc.save(OUT_DOCX)
print(f"已导出冲突清单: {OUT_DOCX}")
